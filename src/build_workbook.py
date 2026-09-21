"""Emit the verification queue as a spreadsheet with a validated label column.

Items already coded in the web tool are carried over. Blind items stay blind: the
rule and zero-shot columns are not written to the sheet.
"""
import glob, json, pathlib, sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = pathlib.Path(__file__).resolve().parents[1]
(ROOT / "annotation").mkdir(exist_ok=True)
CLASSES = ["DEFINITION", "GRANT_OF_AUTHORITY", "SANCTION_REMEDY", "RIGHT_ENTITLEMENT",
           "PROHIBITION", "OBLIGATION", "PROCEDURE", "SCOPE_APPLICABILITY",
           "INSUFFICIENT_CONTEXT"]
GLOSS = {
    "DEFINITION": "Fixes the meaning of a term ('X means...')",
    "GRANT_OF_AUTHORITY": "Confers jurisdiction, power or office on a body",
    "SANCTION_REMEDY": "States a penalty, remedy or enforcement consequence",
    "RIGHT_ENTITLEMENT": "Confers a right, privilege or entitlement on a person",
    "PROHIBITION": "Forbids conduct ('shall not', 'no person may')",
    "OBLIGATION": "Imposes a duty to act, outside a proceeding's steps",
    "PROCEDURE": "Sets the steps, timing or form of a proceeding",
    "SCOPE_APPLICABILITY": "States when, where or to whom the law applies",
    "INSUFFICIENT_CONTEXT": "The excerpt cannot be judged on its own",
}
COLS = [("i", 6), ("label", 21), ("domain", 10), ("heading", 30),
        ("text", 112), ("suggested", 21), ("nation", 26)]

HEAD = PatternFill("solid", fgColor="1F3A4D")
DONE = PatternFill("solid", fgColor="EAF3EC")
THIN = Border(bottom=Side("thin", color="D5DBDF"))


def prior_labels(d):
    """Labels already recorded in the web tool, keyed by queue index."""
    out = {}
    for f in glob.glob(str(pathlib.Path(d) / "*.json")):
        r = json.load(open(f))
        if r.get("label"):
            out[r["i"]] = r["label"]
    return out


def main(done_dir=None):
    items = json.load(open(ROOT / "data/interim/gold_sample.json"))
    done = prior_labels(done_dir) if done_dir else {}

    wb = Workbook()
    ws = wb.active
    ws.title = "code"
    ws.append([c for c, _ in COLS])
    for k, (name, w) in enumerate(COLS, start=1):
        ws.column_dimensions[get_column_letter(k)].width = w
        c = ws.cell(1, k)
        c.fill, c.font = HEAD, Font(bold=True, color="FFFFFF", size=11)
        c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 22

    for it in items:
        ws.append([it["i"], done.get(it["i"], ""), it["domain"], it["heading"],
                   it["text"], it["suggested"] or "", it["nation"]])
        r = ws.max_row
        for k in range(1, len(COLS) + 1):
            cell = ws.cell(r, k)
            cell.border = THIN
            cell.alignment = Alignment(wrap_text=(k in (4, 5)), vertical="top")
            cell.font = Font(size=11)
            if it["i"] in done:
                cell.fill = DONE
        ws.cell(r, 2).font = Font(size=11, bold=True)

    dv = DataValidation(type="list", formula1='"' + ",".join(CLASSES) + '"',
                        allow_blank=True, showDropDown=False)
    dv.error = "Pick one of the nine classes from the dropdown."
    dv.errorTitle = "Not a valid class"
    ws.add_data_validation(dv)
    dv.add(f"B2:B{ws.max_row}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:G{ws.max_row}"

    gs = wb.create_sheet("classes")
    gs.append(["class", "applies when"])
    for k, c in enumerate(CLASSES, start=2):
        gs.append([c, GLOSS[c]])
        gs.cell(k, 1).font = Font(bold=True, size=11)
    for k, w in ((1, 24), (2, 60)):
        gs.column_dimensions[get_column_letter(k)].width = w
    for k in range(1, 3):
        gs.cell(1, k).fill, gs.cell(1, k).font = HEAD, Font(bold=True, color="FFFFFF")

    gs.append([])
    for line in [
        "Label the provision's PRIMARY function, judged on the text as written.",
        "When more than one class fits, the one higher in this list wins.",
        "Do not assess whether the law is good, and do not let the Nation's identity influence the label.",
        "A blank 'suggested' cell means no suggestion was generated. It is not a hint that the text is unusual.",
        "Leave 'label' blank for any row you want to come back to.",
    ]:
        gs.append([line])
    for r in range(len(CLASSES) + 3, gs.max_row + 1):
        gs.cell(r, 1).font = Font(size=11, italic=True)

    out = ROOT / "annotation/provisions_to_code.xlsx"
    wb.save(out)
    print(f"{len(items)} rows ({len(done)} pre-filled) -> {out.relative_to(ROOT)}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else None)
